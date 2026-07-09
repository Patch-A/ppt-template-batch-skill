from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from pptx import Presentation


DEFAULT_MAPPING: dict[str, Any] = {
    "title_shape": 5,
    "buyers_per_slide": 6,
    "slots": [
        {"summary_shape": 15, "products_shape": 23},
        {"summary_group": 16, "summary_child": 2, "products_group": 16, "products_child": 3},
        {"summary_group": 17, "summary_child": 2, "products_group": 17, "products_child": 3},
        {"summary_shape": 19, "products_shape": 26},
        {"summary_shape": 20, "products_shape": 25},
        {"summary_shape": 22, "products_shape": 24},
    ],
}


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def ensure_runs(paragraph, minimum: int) -> None:
    while len(paragraph.runs) < minimum:
        paragraph.add_run()


def get_text_frame(slide, shape_index: int, child_index: int | None = None):
    shape = slide.shapes[shape_index - 1]
    if child_index is not None:
        shape = shape.shapes[child_index - 1]
    return shape.text_frame


def set_single_run_text(text_frame, value: str) -> None:
    paragraph = text_frame.paragraphs[0]
    ensure_runs(paragraph, 1)
    paragraph.runs[0].text = value
    for run in paragraph.runs[1:]:
        run.text = ""


def split_summary(name: str, summary: str) -> tuple[str, str]:
    if summary.startswith(name):
        return name, summary[len(name) :]
    return name, summary


def set_summary_text(text_frame, buyer: dict[str, Any]) -> None:
    name = str(buyer.get("name", "") or "")
    summary = str(buyer.get("summary") or buyer.get("intro") or "")
    paragraph = text_frame.paragraphs[0]
    ensure_runs(paragraph, 2)
    name_text, body_text = split_summary(name, summary)
    paragraph.runs[0].text = name_text
    paragraph.runs[1].text = body_text
    for run in paragraph.runs[2:]:
        run.text = ""


def products_text(buyer: dict[str, Any]) -> str:
    value = str(buyer.get("products") or buyer.get("categories") or "").strip()
    if value.startswith("采购品类："):
        return value
    return f"采购品类：{value}" if value else ""


def fill_slot(slide, slot: dict[str, Any], buyer: dict[str, Any]) -> None:
    if "summary_group" in slot:
        summary_frame = get_text_frame(slide, int(slot["summary_group"]), int(slot["summary_child"]))
    else:
        summary_frame = get_text_frame(slide, int(slot["summary_shape"]))
    set_summary_text(summary_frame, buyer)

    if "products_group" in slot:
        products_frame = get_text_frame(slide, int(slot["products_group"]), int(slot["products_child"]))
    else:
        products_frame = get_text_frame(slide, int(slot["products_shape"]))
    set_single_run_text(products_frame, products_text(buyer))


def fill_slide(slide, page: dict[str, Any], mapping: dict[str, Any]) -> None:
    buyers = page["buyers"]
    buyers_per_slide = int(mapping.get("buyers_per_slide", len(mapping["slots"])))
    if len(buyers) != buyers_per_slide:
        raise ValueError(f"Each briefing slide must contain exactly {buyers_per_slide} buyers.")

    title_frame = get_text_frame(slide, int(mapping["title_shape"]))
    set_single_run_text(title_frame, str(page["title"]))

    for slot, buyer in zip(mapping["slots"], buyers):
        fill_slot(slide, slot, buyer)


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Fill buyer-briefing PPT pages while preserving template run-level text styles."
    )
    parser.add_argument("template", help="Template PPTX path")
    parser.add_argument("pages_json", help="JSON with pages: title + 6 buyers per page")
    parser.add_argument("output", help="Output PPTX path")
    parser.add_argument("--layout-config", help="Optional briefing layout mapping JSON")
    args = parser.parse_args()

    pages = load_json(Path(args.pages_json))
    mapping = DEFAULT_MAPPING
    if args.layout_config:
        mapping = load_json(Path(args.layout_config))

    prs = Presentation(args.template)
    if len(pages) > len(prs.slides):
        raise ValueError("Template does not have enough slides. Duplicate the briefing page layout first.")

    for index, page in enumerate(pages):
        fill_slide(prs.slides[index], page, mapping)

    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
